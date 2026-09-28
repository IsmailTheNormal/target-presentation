import openpyxl

# 1. English Only Master Schedule
wb = openpyxl.load_workbook('/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx')
ws = wb['Umumiy Jadval']

fixed_breaks = ["NONUSHTA", "TUSHLIK", "POLNIK"]

for row in range(2, 67):
    for col in range(6, 26):
        col_letter = openpyxl.utils.get_column_letter(col)
        val = ws[f"{col_letter}{row}"].value
        if val:
            if not ("English" in val or "ELT" in val or any(b in val for b in fixed_breaks)):
                ws[f"{col_letter}{row}"] = ""

for sheet in wb.sheetnames[1:]:
    ws_class = wb[sheet]
    for row in range(2, 67):
        val = ws_class[f"D{row}"].value
        if val:
            if not ("English" in val or "ELT" in val or any(b in val for b in fixed_breaks)):
                ws_class[f"D{row}"] = ""

wb.save('/home/dpdp/target/timetable/ENGLISH_ONLY_MASTER.xlsx')

# 2. English Only Teacher POV
wb_t = openpyxl.load_workbook('/home/dpdp/target/timetable/TEACHER_POV_JADVAL.xlsx')
ws_t = wb_t.active

# Columns 4 to end are teachers.
# We want to keep KUN(1), SOAT(2), VAQT(3), and ELT teachers + Primary teachers
cols_to_delete = []
for col in range(4, ws_t.max_column + 1):
    header = ws_t.cell(row=1, column=col).value
    if header:
        if not ("ELT" in header or "1-" in header or "2-" in header or "3-" in header or "4-" in header):
            cols_to_delete.append(col)

# Delete from back to front to avoid index shifting issues
for col in sorted(cols_to_delete, reverse=True):
    ws_t.delete_cols(col)

# Also clear out any non-English classes from the primary teachers (since they teach everything)
for row in range(2, 67):
    for col in range(4, ws_t.max_column + 1):
        val = ws_t.cell(row=row, column=col).value
        if val and not ("English" in val or "ELT" in val):
            # For primary teachers, the string is like "1-A Math". We only keep it if it has English.
            if any(p in ws_t.cell(row=1, column=col).value for p in ["1-", "2-", "3-", "4-"]):
                if "English" not in val:
                    ws_t.cell(row=row, column=col).value = ""
            
wb_t.save('/home/dpdp/target/timetable/ENGLISH_ONLY_TEACHER_POV.xlsx')
print("Created English Only files!")
