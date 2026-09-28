import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

classes_1_4 = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]
classes_5_11 = ["5-A", "5-B", "6-A", "6-B", "7-A", "7-B", "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
all_classes = classes_1_4 + classes_5_11

# Days and indices (13 rows per day)
# We have 5 days * 13 rows = 65 rows (plus 1 header row)
# Real rows are 2 to 66.

def get_empty_slots(col_letter):
    slots = []
    for row in range(2, 67):
        val = ws[f"{col_letter}{row}"].value
        # Ignore Lunch, Polnik, Togarak, etc.
        if not val:
            slots.append(row)
    return slots

def find_2_hour_blocks(slots):
    blocks = []
    i = 0
    while i < len(slots)-1:
        if slots[i+1] == slots[i] + 1:
            # check if they are in the same day (not crossing day boundary)
            # Row 2 to 14 is Mon, 15 to 27 is Tue...
            day1 = (slots[i]-2) // 13
            day2 = (slots[i+1]-2) // 13
            if day1 == day2:
                blocks.append((slots[i], slots[i+1]))
                i += 2
                continue
        i += 1
    return blocks

# Requirements per class
for col_idx, c in enumerate(all_classes):
    col_letter = openpyxl.utils.get_column_letter(col_idx + 5) # E is KUN, SOAT... actually KUN=A, SOAT1=B, VAQT1=C, SOAT2=D, VAQT2=E. Class starts at F (which is 6)
    # Wait, let's check columns.
    # A=KUN, B=SOAT(1-4), C=VAQT(1-4), D=SOAT(5-11), E=VAQT(5-11)
    # So 1-A is F (6), 1-B is G (7), etc.
    col_letter = openpyxl.utils.get_column_letter(col_idx + 6)
    
    slots = get_empty_slots(col_letter)
    blocks_2h = find_2_hour_blocks(slots)
    
    # 1. English & Russian
    if c.endswith("A"):
        eng_blocks = 4 # 8 hours
        rus_blocks = 2; rus_single = 1 # 5 hours
    else:
        rus_blocks = 4 # 8 hours
        eng_blocks = 2; eng_single = 1 # 5 hours
        
    # 2. Math (assume 3 blocks of 2 = 6 hours)
    math_blocks = 3
    
    # 3. IT (1 block of 2 = 2 hours)
    it_blocks = 1
    
    # Place blocks
    used_rows = set()
    
    def place_block(subject, count):
        placed = 0
        for block in blocks_2h:
            if placed >= count: break
            if block[0] not in used_rows and block[1] not in used_rows:
                ws[f"{col_letter}{block[0]}"] = subject
                ws[f"{col_letter}{block[1]}"] = subject
                used_rows.add(block[0])
                used_rows.add(block[1])
                placed += 1
                
    def place_single(subject, count):
        placed = 0
        for r in slots:
            if placed >= count: break
            if r not in used_rows:
                ws[f"{col_letter}{r}"] = subject
                used_rows.add(r)
                placed += 1

    place_block("English", eng_blocks)
    place_block("Russian", rus_blocks)
    place_block("Math", math_blocks)
    place_block("IT", it_blocks)
    
    if c.endswith("A"):
        place_single("Russian", rus_single)
        place_single("English", 0)
    else:
        place_single("English", eng_single)
        place_single("Russian", 0)
        
    # Place Uzbek (3 single hours)
    place_single("Uzbek Language", 3)
    
    # Place Geo, Bio, Chem (1 hour each) for grades > 4
    if not c.startswith(("1-", "2-", "3-", "4-")):
        place_single("Geography", 1)
        place_single("Biology", 1)
        place_single("Chemistry", 1)

wb.save(file_path)
print("Subjects filled!")
