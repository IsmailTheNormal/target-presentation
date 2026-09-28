import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

classes_1_4 = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]
classes_5_11 = ["5-A", "5-B", "6-A", "6-B", "7-A", "7-B", "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
all_classes = classes_1_4 + classes_5_11

def get_teacher(cls_name, subject):
    # Exact mappings
    if cls_name == "1-A" and "English" in subject: return "Baxritdinova Gulya"
    if cls_name == "1-B" and "Russian" in subject: return "Fozilova Nargiza"
    if cls_name == "2-A" and "English" in subject: return "Lucille Johnalyn"
    if cls_name == "2-B" and "Russian" in subject: return "Xosiyat"
    if cls_name == "3-A" and "English" in subject: return "Adilova Nigora"
    if cls_name == "4-A" and "English" in subject: return "Nida"
    if cls_name == "5-A" and "English" in subject: return "Robiya Ilyasovna"
    
    # Subject generic mappings based on Dars taqsimoti
    if "Uzbek" in subject: return "Farangiz"
    if "Mental arifmetika" in subject: return "Ra'no"
    if "Science" in subject: return "Durdona"
    if "Geography" in subject or "History" in subject: return "Diyor Orifovich"
    if "Biology" in subject or "Chemistry" in subject: return "Zufar"
    if "IT" in subject: return "Mamarajab/Ismoil"
    if "Math" in subject: return "Shoxsanam/Lutfullo"
    if "English" in subject: return "Warner/Firdavs/Surayyo"
    if "Russian" in subject: return "Saodat Maqsudovna"
    if "Business" in subject or "economic" in subject: return "Muhammadali"
    
    return ""

for col_idx, c in enumerate(all_classes):
    col_letter = openpyxl.utils.get_column_letter(col_idx + 6)
    ws_class = wb[c]
    
    for row in range(2, 67):
        val = ws[f"{col_letter}{row}"].value
        if val and val not in ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak (English/Math Uzum Market)"]:
            teacher = get_teacher(c, val)
            if teacher:
                new_val = f"{val}\n({teacher})"
                ws[f"{col_letter}{row}"] = new_val
                ws_class[f"D{row}"] = new_val

wb.save(file_path)
print("Teachers added!")
