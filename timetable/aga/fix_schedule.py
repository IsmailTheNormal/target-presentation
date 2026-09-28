import pandas as pd
import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

cohorts = [
    ["5-A", "5-B"],
    ["6-A", "6-B", "7-A", "7-B"],
    ["8-A", "8-B", "9-A", "9-B"],
    ["10-A", "10-B", "11-A", "11-B"]
]
classes_1_4 = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]
all_classes = classes_1_4 + [c for cohort in cohorts for c in cohort]

def get_col(c):
    return openpyxl.utils.get_column_letter(all_classes.index(c) + 6)

fixed_vals = ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak", 
              "Jismoniy tarbiya", "Science", "Shaxmat", "Mental arifmetika"]

# Clear all non-fixed slots
for row in range(2, 67):
    for c in all_classes:
        col = get_col(c)
        val = ws[f"{col}{row}"].value
        if val:
            is_fixed = any(f in val for f in fixed_vals)
            if not is_fixed:
                ws[f"{col}{row}"] = None

def get_empty_slots(c_list):
    slots = []
    for row in range(2, 67):
        is_empty = all(ws[f"{get_col(c)}{row}"].value is None for c in c_list)
        if is_empty:
            slots.append(row)
    return slots

def find_2_hour_blocks(slots):
    blocks = []
    i = 0
    while i < len(slots)-1:
        if slots[i+1] == slots[i] + 1:
            day1 = (slots[i]-2) // 13
            day2 = (slots[i+1]-2) // 13
            if day1 == day2:
                blocks.append((slots[i], slots[i+1]))
                i += 2
                continue
        i += 1
    return blocks

# Fill cohort synchronized subjects (5-11)
for cohort in cohorts:
    slots = get_empty_slots(cohort)
    blocks = find_2_hour_blocks(slots)
    used_rows = set()
    used_days = {'Math': set(), 'ELT': set()}
    
    placed_math = 0
    for block in blocks:
        if placed_math >= 3: break
        day = (block[0]-2) // 13
        if block[0] not in used_rows and block[1] not in used_rows and day not in used_days['Math']:
            for c in cohort:
                ws[f"{get_col(c)}{block[0]}"] = "Math (Levels)"
                ws[f"{get_col(c)}{block[1]}"] = "Math (Levels)"
            used_rows.add(block[0]); used_rows.add(block[1])
            used_days['Math'].add(day)
            placed_math += 1
            
    placed_eng = 0
    for block in blocks:
        if placed_eng >= 2: break
        day = (block[0]-2) // 13
        if block[0] not in used_rows and block[1] not in used_rows and day not in used_days['ELT']:
            for c in cohort:
                ws[f"{get_col(c)}{block[0]}"] = "ELT (Levels)"
                ws[f"{get_col(c)}{block[1]}"] = "ELT (Levels)"
            used_rows.add(block[0]); used_rows.add(block[1])
            used_days['ELT'].add(day)
            placed_eng += 1

# Fill independent subjects
for c in all_classes:
    col = get_col(c)
    slots = get_empty_slots([c])
    blocks = find_2_hour_blocks(slots)
    used_rows = set()
    used_days = {'English': set(), 'Russian': set(), 'IT': set(), 'Math': set()}
    
    is_1_4 = c in classes_1_4
    
    eng_rem_blocks = 4 if is_1_4 and c.endswith("A") else (2 if c.endswith("A") else 0)
    eng_rem_single = 0 if c.endswith("A") else 1
    
    rus_rem_blocks = 2 if is_1_4 and c.endswith("A") else (4 if c.endswith("B") else 1)
    rus_rem_single = 1 if c.endswith("A") else 0
    
    math_rem_blocks = 3 if is_1_4 else 0 # 1-4 Math
    
    def place_blocks(subject, limit):
        placed = 0
        for block in blocks:
            if placed >= limit: break
            day = (block[0]-2) // 13
            if block[0] not in used_rows and block[1] not in used_rows and day not in used_days[subject]:
                label = subject
                if is_1_4: # Use main teacher name instead of "ORS"
                    t = {"1-A":"Baxritdinova Gulya", "1-B":"Fozilova Nargiza", "2-A":"Lucille Johnalyn", "2-B":"Xosiyat", "3-A":"Adilova Nigora", "4-A":"Nida"}[c]
                    label = f"{subject}\n({t})"
                
                ws[f"{col}{block[0]}"] = label
                ws[f"{col}{block[1]}"] = label
                used_rows.add(block[0]); used_rows.add(block[1])
                used_days[subject].add(day)
                placed += 1
                
    place_blocks('English', eng_rem_blocks)
    place_blocks('Russian', rus_rem_blocks)
    place_blocks('Math', math_rem_blocks)
    place_blocks('IT', 1)

    # Singles
    singles = ["Uzbek Language"]*3
    if not is_1_4: singles += ["Geography", "Biology", "Chemistry"]
    if eng_rem_single: singles.append("English")
    if rus_rem_single: singles.append("Russian")
    
    idx = 0
    for r in get_empty_slots([c]):
        if idx >= len(singles): break
        if r not in used_rows:
            label = singles[idx]
            if is_1_4 and label in ["Uzbek Language", "English", "Russian"]:
                t = {"1-A":"Baxritdinova Gulya", "1-B":"Fozilova Nargiza", "2-A":"Lucille Johnalyn", "2-B":"Xosiyat", "3-A":"Adilova Nigora", "4-A":"Nida"}[c]
                label = f"{label}\n({t})"
            ws[f"{col}{r}"] = label
            used_rows.add(r)
            idx += 1

# Sync to POV
for c in all_classes:
    col = get_col(c)
    ws_class = wb[c]
    for row in range(2, 67):
        val = ws[f"{col}{row}"].value
        ws_class[f"D{row}"] = val if val else ""

wb.save(file_path)
print("Schedule fixed: No 4-hour days, no ORs!")
