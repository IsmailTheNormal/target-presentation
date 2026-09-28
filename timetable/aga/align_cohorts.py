import pandas as pd
import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

cohorts = [
    ["5-A", "5-B"],
    ["6-A", "6-B", "7-A", "7-B"],
    ["8-A", "8-B", "9-A", "9-B"],
    ["10-A", "10-B", "11-A", "11-B"]
]

# Map class name to column index (E is 5, F is 6 -> 1-A)
def get_col(c):
    all_c = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
             "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
             "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
    return openpyxl.utils.get_column_letter(all_c.index(c) + 6)

fixed_vals = ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak (English/Math Uzum Market)", 
              "Jismoniy tarbiya (combined)", "Science", "Shaxmat", "Mental arifmetika"]

# Clear 5-11
for row in range(2, 67):
    for c in [c for cohort in cohorts for c in cohort]:
        col = get_col(c)
        val = ws[f"{col}{row}"].value
        if val:
            # check if it's fixed
            is_fixed = False
            for f in fixed_vals:
                if f in val: is_fixed = True
            if not is_fixed:
                ws[f"{col}{row}"] = None

def get_shared_empty_slots(cohort):
    slots = []
    for row in range(2, 67):
        is_empty = True
        for c in cohort:
            if ws[f"{get_col(c)}{row}"].value is not None:
                is_empty = False
                break
        if is_empty:
            slots.append(row)
    return slots

def find_2_hour_blocks(slots):
    blocks = []
    i = 0
    while i < len(slots)-1:
        if slots[i+1] == slots[i] + 1:
            day1 = (slots[i]-2) // 13
            day2 = (slots[i+1]-2) // 13
            if day1 == day2:
                blocks.append((slots[i], slots[i+1]))
                i += 2
                continue
        i += 1
    return blocks

# Fill cohort synchronized subjects
for cohort in cohorts:
    slots = get_shared_empty_slots(cohort)
    blocks = find_2_hour_blocks(slots)
    
    used_rows = set()
    
    # We synchronize 2 blocks (4 hours) of English, and 3 blocks (6 hours) of Math
    # Wait, B classes only have 5 hours of English. So we can sync 2 blocks (4 hours) of English for everyone.
    # The remaining 1 or 4 hours of English can be placed independently later.
    
    # Place Math (3 blocks = 6 hours synced)
    placed_math = 0
    for block in blocks:
        if placed_math >= 3: break
        if block[0] not in used_rows and block[1] not in used_rows:
            for c in cohort:
                ws[f"{get_col(c)}{block[0]}"] = "Math (Levels A/B/C)"
                ws[f"{get_col(c)}{block[1]}"] = "Math (Levels A/B/C)"
            used_rows.add(block[0])
            used_rows.add(block[1])
            placed_math += 1
            
    # Place English (2 blocks = 4 hours synced)
    placed_eng = 0
    for block in blocks:
        if placed_eng >= 2: break
        if block[0] not in used_rows and block[1] not in used_rows:
            for c in cohort:
                label = "ELT (Ordinary/Advanced)" if cohort == cohorts[0] else "ELT (Levels)"
                ws[f"{get_col(c)}{block[0]}"] = label
                ws[f"{get_col(c)}{block[1]}"] = label
            used_rows.add(block[0])
            used_rows.add(block[1])
            placed_eng += 1

# Now fill the remaining stuff independently for each class
for cohort in cohorts:
    for c in cohort:
        col = get_col(c)
        # get empty slots for this specific class
        slots = []
        for row in range(2, 67):
            if not ws[f"{col}{row}"].value:
                slots.append(row)
                
        blocks = find_2_hour_blocks(slots)
        used = set()
        
        eng_rem_blocks = 2 if c.endswith("A") else 0
        eng_rem_single = 0 if c.endswith("A") else 1
        
        rus_rem_blocks = 1 if c.endswith("A") else 4
        rus_rem_single = 1 if c.endswith("A") else 0
        
        # Place English remaining
        placed = 0
        for block in blocks:
            if placed >= eng_rem_blocks: break
            if block[0] not in used and block[1] not in used:
                ws[f"{col}{block[0]}"] = "English"
                ws[f"{col}{block[1]}"] = "English"
                used.add(block[0]); used.add(block[1])
                placed += 1
                
        # Place Russian remaining
        placed = 0
        for block in blocks:
            if placed >= rus_rem_blocks: break
            if block[0] not in used and block[1] not in used:
                ws[f"{col}{block[0]}"] = "Russian"
                ws[f"{col}{block[1]}"] = "Russian"
                used.add(block[0]); used.add(block[1])
                placed += 1
                
        # IT
        placed = 0
        for block in blocks:
            if placed >= 1: break
            if block[0] not in used and block[1] not in used:
                ws[f"{col}{block[0]}"] = "IT"
                ws[f"{col}{block[1]}"] = "IT"
                used.add(block[0]); used.add(block[1])
                placed += 1
                
        # Singles
        singles_to_place = ["Uzbek Language"]*3 + ["Geography", "Biology", "Chemistry"]
        if eng_rem_single: singles_to_place.append("English")
        if rus_rem_single: singles_to_place.append("Russian")
        
        idx = 0
        for r in slots:
            if idx >= len(singles_to_place): break
            if r not in used:
                ws[f"{col}{r}"] = singles_to_place[idx]
                used.add(r)
                idx += 1

# Sync to POV sheets
all_c = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
         "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
         "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
for c in all_c:
    col = get_col(c)
    ws_class = wb[c]
    for row in range(2, 67):
        val = ws[f"{col}{row}"].value
        if val:
            ws_class[f"D{row}"] = val
        else:
            ws_class[f"D{row}"] = ""

wb.save(file_path)
print("Cohorts aligned!")
