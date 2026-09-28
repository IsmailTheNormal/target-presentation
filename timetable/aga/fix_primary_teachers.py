import openpyxl
import re

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

main_teachers = {
    "1-A": "Baxritdinova Gulya",
    "1-B": "Fozilova Nargiza",
    "2-A": "Lucille Johnalyn",
    "2-B": "Xosiyat",
    "3-A": "Adilova Nigora",
    "4-A": "Nida"
}

classes_1_4 = list(main_teachers.keys())

for col_idx, c in enumerate(wb.sheetnames[1:]): # skip Umumiy Jadval
    if c not in classes_1_4:
        continue
    
    col_letter = openpyxl.utils.get_column_letter(classes_1_4.index(c) + 6)
    ws_class = wb[c]
    main_teacher = main_teachers[c]
    
    for row in range(2, 67):
        val = ws[f"{col_letter}{row}"].value
        if val and "\n(" in str(val):
            # We don't overwrite specials if they have their own teacher, but user said "all classes". 
            # Let's overwrite English, Russian, Math, Uzbek, IT, Science, etc.
            # Actually, the easiest is to just swap out the name inside the parenthesis for these subjects.
            subject = val.split("\n(")[0]
            if subject not in ["Mental arifmetika", "Jismoniy tarbiya (combined)"]:
                new_val = f"{subject}\n({main_teacher})"
                ws[f"{col_letter}{row}"] = new_val
                ws_class[f"D{row}"] = new_val

wb.save(file_path)
print("Primary teachers fixed!")
