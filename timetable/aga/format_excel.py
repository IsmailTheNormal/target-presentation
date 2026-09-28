import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'), 
    top=Side(style='thin'), bottom=Side(style='thin')
)
center_aligned = Alignment(horizontal='center', vertical='center', wrap_text=True)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # 1. Adjust column widths
    for col in ws.columns:
        column_letter = col[0].column_letter
        max_length = 0
        for cell in col:
            if cell.value:
                # Calculate approximate length
                length = len(str(cell.value))
                if length > max_length:
                    max_length = length
        
        # Cap width at 25 so it's readable but wraps long text
        adjusted_width = min(max_length + 3, 25)
        # Ensure time columns are wide enough
        if sheet_name == 'Umumiy Jadval' and column_letter in ['A', 'B', 'C', 'D', 'E']:
            adjusted_width = max(15, adjusted_width)
            
        ws.column_dimensions[column_letter].width = adjusted_width
        
    # 2. Format cells (borders, alignment, header)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = center_aligned
            cell.border = thin_border
            
            # Format header row
            if cell.row == 1:
                cell.font = header_font
                cell.fill = header_fill

wb.save(file_path)
print("Formatting complete!")
