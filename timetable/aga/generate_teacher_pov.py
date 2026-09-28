import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from collections import defaultdict

wb_master = openpyxl.load_workbook('/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx')
ws_master = wb_master['Umumiy Jadval']

all_classes = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
               "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
               "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]

teacher_mapping = {
    "1-A": "Baxritdinova Gulya", "1-B": "Fozilova Nargiza",
    "2-A": "Lucille Johnalyn", "2-B": "Xosiyat",
    "3-A": "Adilova Nigora", "4-A": "Nida"
}

teachers = list(teacher_mapping.values()) + [
    "Ra'no (Mental Arifmetika)", 
    "Firdavs (ELT)", "Warner/Werner (ELT)", "Surayyo (ELT)", "Sabrina (ELT)", "Robiya Ilyasovna (ELT)",
    "Shoxsanam (Math)", "Lutfullo (Math)", "Ilyos (Math)",
    "Mamarajab (IT)", "Ismoil (IT)", "Nigmatov (IT)", "Nigina (IT)",
    "Farangiz (Uzbek)", "Diyor Orifovich (Geo/Hist)", "Durdona (Science)",
    "Zufar (Bio/Chem)", "Saodat Maqsudovna (Russian)"
]

def assign_teachers(c, val):
    if not val: return []
    if any(b in val for b in ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak", "Jismoniy tarbiya", "Shaxmat"]):
        return []
        
    res = []
    if c in teacher_mapping:
        if "Mental arifmetika" in val:
            res.append(("Ra'no (Mental Arifmetika)", f"{c}"))
        else:
            res.append((teacher_mapping[c], f"{c} {val}"))
    else:
        if "ELT" in val or "English" in val:
            if c.startswith(("6-", "7-")): res.extend([("Firdavs (ELT)", c), ("Warner/Werner (ELT)", c), ("Surayyo (ELT)", c)])
            elif c.startswith(("8-", "9-")): res.extend([("Sabrina (ELT)", c), ("Surayyo (ELT)", c), ("Warner/Werner (ELT)", c)])
            elif c.startswith(("10-", "11-")): res.extend([("Warner/Werner (ELT)", c), ("Firdavs (ELT)", c), ("Surayyo (ELT)", c)])
            elif c.startswith("5-"): res.append(("Robiya Ilyasovna (ELT)", c))
        elif "Math" in val:
            res.extend([("Shoxsanam (Math)", c), ("Lutfullo (Math)", c), ("Ilyos (Math)", c)])
        elif "IT" in val:
            # IT has 4 teachers, maybe 2 pairs. We'll just list them all for IT slots
            res.extend([("Mamarajab (IT)", c), ("Ismoil (IT)", c), ("Nigmatov (IT)", c), ("Nigina (IT)", c)])
        elif "Uzbek" in val: res.append(("Farangiz (Uzbek)", f"{c}"))
        elif "Geography" in val or "History" in val: res.append(("Diyor Orifovich (Geo/Hist)", f"{c}"))
        elif "Science" in val: res.append(("Durdona (Science)", f"{c}"))
        elif "Biology" in val or "Chemistry" in val: res.append(("Zufar (Bio/Chem)", f"{c}"))
        elif "Russian" in val: res.append(("Saodat Maqsudovna (Russian)", f"{c}"))
        
    return res

teacher_grid = defaultdict(lambda: defaultdict(list))

for row in range(2, 67):
    for idx, c in enumerate(all_classes):
        col_letter = openpyxl.utils.get_column_letter(idx + 6)
        val = ws_master[f"{col_letter}{row}"].value
        
        assigned = assign_teachers(c, val)
        for t, label in assigned:
            if label not in teacher_grid[row][t]:
                teacher_grid[row][t].append(label)

wb_new = openpyxl.Workbook()
ws_new = wb_new.active
ws_new.title = "O'qituvchilar POV"

ws_new["A1"] = "KUN"
ws_new["B1"] = "SOAT"
ws_new["C1"] = "VAQT"

for idx, t in enumerate(teachers):
    ws_new.cell(row=1, column=idx+4, value=t)

for row in range(2, 67):
    kun = ws_master[f"A{row}"].value
    soat = ws_master[f"D{row}"].value
    vaqt = ws_master[f"E{row}"].value
    
    ws_new.cell(row=row, column=1, value=kun if kun else "")
    ws_new.cell(row=row, column=2, value=soat if soat else "")
    ws_new.cell(row=row, column=3, value=vaqt if vaqt else "")
    
    for idx, t in enumerate(teachers):
        classes_taught = teacher_grid[row][t]
        if classes_taught:
            classes_taught.sort()
            ws_new.cell(row=row, column=idx+4, value=", ".join(classes_taught))

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
center_aligned = Alignment(horizontal='center', vertical='center', wrap_text=True)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")

for col in ws_new.columns:
    column_letter = col[0].column_letter
    ws_new.column_dimensions[column_letter].width = 18
    
for row in ws_new.iter_rows():
    for cell in row:
        cell.alignment = center_aligned
        cell.border = thin_border
        if cell.row == 1:
            cell.font = header_font
            cell.fill = header_fill

out_path = '/home/dpdp/target/timetable/TEACHER_POV_JADVAL.xlsx'
wb_new.save(out_path)
print("Teacher POV updated with ALL individual teachers!")
