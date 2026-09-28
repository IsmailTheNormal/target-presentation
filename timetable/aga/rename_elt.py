import openpyxl

files = ['/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx', '/home/dpdp/target/timetable/TEACHER_POV_JADVAL.xlsx']

for f in files:
    wb = openpyxl.load_workbook(f)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if "ELT" in cell.value:
                        cell.value = cell.value.replace("ELT", "English")
    wb.save(f)
print("Replaced ELT with English in all files!")
