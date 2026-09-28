import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

all_classes = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
               "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
               "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]

fillers_5_11 = ["Robotics", "Financial Literacy", "Art", "Music", "Personal Development", "Mindlab (Werner)", "Project Work"]
fillers_1_4 = ["Robotics", "Financial Literacy", "Art", "Music", "Personal Development", "Reading", "Mindlab"]

for c in all_classes:
    col = openpyxl.utils.get_column_letter(all_classes.index(c) + 6)
    filler_idx = 0
    ws_class = wb[c]
    is_1_4 = c in ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]
    fillers = fillers_1_4 if is_1_4 else fillers_5_11
    
    for r in range(2, 67):
        val = ws[f"{col}{r}"].value
        if val is None:
            # fill the gap
            sub = fillers[filler_idx % len(fillers)]
            ws[f"{col}{r}"] = sub
            ws_class[f"D{r}"] = sub
            filler_idx += 1

wb.save(file_path)
print("Missing subjects filled! Grid is 100% packed.")
