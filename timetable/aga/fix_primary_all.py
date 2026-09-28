import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

main_teachers = {
    "1-A": "Baxritdinova Gulya",
    "1-B": "Fozilova Nargiza",
    "2-A": "Lucille Johnalyn",
    "2-B": "Xosiyat",
    "3-A": "Adilova Nigora",
    "4-A": "Nida"
}

classes_1_4 = list(main_teachers.keys())
fixed_breaks = ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak"]

for c in classes_1_4:
    col_letter = openpyxl.utils.get_column_letter(classes_1_4.index(c) + 6)
    ws_class = wb[c]
    teacher = main_teachers[c]
    
    for row in range(2, 67):
        val = ws[f"{col_letter}{row}"].value
        if val:
            # Skip breaks/lunches/togarak
            is_break = any(b in val for b in fixed_breaks)
            if not is_break:
                # Clean up the subject name (remove any existing parenthesis)
                subject = val.split("\n(")[0].strip()
                subject = subject.replace(" (combined)", "")
                
                new_val = f"{subject}\n({teacher})"
                ws[f"{col_letter}{row}"] = new_val
                ws_class[f"D{row}"] = new_val

wb.save(file_path)
print("All 1-4 classes updated to primary teacher!")
