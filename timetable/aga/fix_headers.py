import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

main_teachers = {
    "1-A": "Baxritdinova Gulya",
    "1-B": "Fozilova Nargiza",
    "2-A": "Lucille Johnalyn",
    "2-B": "Xosiyat",
    "3-A": "Adilova Nigora",
    "4-A": "Nida"
}

classes_1_4 = list(main_teachers.keys())

for c in classes_1_4:
    # 1-A is F (6), 1-B is G (7), etc.
    col_idx = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
               "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
               "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"].index(c) + 6
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    
    ws_class = wb[c]
    teacher = main_teachers[c]
    
    # Update Master Header
    ws[f"{col_letter}1"] = f"{c}\n({teacher})"
    
    # Strip from cells
    for row in range(2, 67):
        val = ws[f"{col_letter}{row}"].value
        if val and "\n(" in str(val):
            subject = val.split("\n(")[0].strip()
            ws[f"{col_letter}{row}"] = subject
            ws_class[f"D{row}"] = subject

wb.save(file_path)
print("Headers fixed!")
