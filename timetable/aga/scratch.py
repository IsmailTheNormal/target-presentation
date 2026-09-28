import pandas as pd
import os
import glob
try:
    from docx import Document
except ImportError:
    os.system('pip install python-docx pandas openpyxl')
    from docx import Document

dir_path = '/home/dpdp/target/timetable'
files = sorted(glob.glob(os.path.join(dir_path, '*')))

for f in files:
    if f.endswith('.xlsx') and not os.path.basename(f).startswith('~'):
        print(f"=== Excel File: {os.path.basename(f)} ===")
        try:
            xl = pd.ExcelFile(f)
            print("Sheets:", xl.sheet_names)
            for sheet in xl.sheet_names:
                df = xl.parse(sheet)
                print(f" Sheet '{sheet}' shape: {df.shape}")
                print(df.head(10).to_string())
        except Exception as e:
            print("Error reading excel:", e)
        print("\n")
    elif f.endswith('.docx') and not os.path.basename(f).startswith('~'):
        print(f"=== Word File: {os.path.basename(f)} ===")
        try:
            doc = Document(f)
            print(f"Paragraphs count: {len(doc.paragraphs)}")
            for i, p in enumerate(doc.paragraphs[:15]):
                if p.text.strip():
                    print(f" {p.text.strip()}")
            print(f"Tables count: {len(doc.tables)}")
            if len(doc.tables) > 0:
                print("Table 1 dimensions: {} rows, {} cols".format(len(doc.tables[0].rows), len(doc.tables[0].columns)))
                for row in doc.tables[0].rows[:3]:
                    print("  |  ".join([cell.text.strip().replace('\n', ' ') for cell in row.cells]))
        except Exception as e:
            print("Error reading word:", e)
        print("\n")
