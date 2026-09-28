import openpyxl

wb = openpyxl.load_workbook('/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx')
ws = wb['Umumiy Jadval']

print("--- Monday 6-A ---")
for r in range(2, 15):
    soat = ws[f"D{r}"].value
    vaqt = ws[f"E{r}"].value
    val = ws[f"M{r}"].value # M is 6-A
    print(f"Row {r} | {soat} ({vaqt}): {val}")

print("\n--- Monday 1-A ---")
for r in range(2, 15):
    soat = ws[f"B{r}"].value
    vaqt = ws[f"C{r}"].value
    val = ws[f"F{r}"].value # F is 1-A
    print(f"Row {r} | {soat} ({vaqt}): {val}")
