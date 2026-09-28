import openpyxl

wb = openpyxl.load_workbook('/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx')
ws = wb['Umumiy Jadval']

print("--- Friday 6-A ---")
for r in range(54, 67):
    soat = ws[f"D{r}"].value
    val = ws[f"M{r}"].value # M is 6-A
    print(f"Row {r} | {soat}: {val}")
