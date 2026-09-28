import openpyxl
import random

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

all_classes = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
               "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
               "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
def get_col(c): return openpyxl.utils.get_column_letter(all_classes.index(c) + 6)

def get_teacher(c, sub):
    if not sub: return None
    if "(Levels)" in sub or "(Ordinary/Advanced)" in sub or "To'garak" in sub or "(combined)" in sub or "Shaxmat" in sub: return None 
    if sub in ["NONUSHTA", "TUSHLIK", "POLNIK"]: return None
    if c in ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]:
        if sub in ["English", "Russian", "Math", "Uzbek Language"]: return f"Primary_{c}"
    if "Russian" in sub and c not in ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A"]: return "Saodat"
    if "Uzbek" in sub: return "Farangiz"
    if "Science" in sub: return "Durdona"
    if "Geography" in sub or "History" in sub: return "Diyor"
    if "Biology" in sub or "Chemistry" in sub: return "Zufar"
    if "Business" in sub or "Financial" in sub or "Economic" in sub: return "Muhammadali"
    if "IT" in sub: return "IT_Dept"
    return sub 

def is_academic(c, r):
    val = ws[f"{get_col(c)}{r}"].value
    if not val: return False
    if any(b in val for b in ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak"]): return False
    if "(Levels)" in val or "(Ordinary/Advanced)" in val or "(combined)" in val or "Shaxmat" in val: return False
    return True

for _ in range(50):
    clash_found = False
    for r in range(2, 67):
        t_counts = {}
        for c in all_classes:
            if is_academic(c, r):
                t = get_teacher(c, ws[f"{get_col(c)}{r}"].value)
                if t: t_counts[t] = t_counts.get(t, 0) + 1
                    
        for c in all_classes:
            if is_academic(c, r):
                val = ws[f"{get_col(c)}{r}"].value
                t = get_teacher(c, val)
                if not t: continue
                
                cap = 2 if t == "IT_Dept" else 1
                if t_counts[t] > cap:
                    clash_found = True
                    for r2 in range(2, 67):
                        if r == r2 or not is_academic(c, r2): continue
                        val2 = ws[f"{get_col(c)}{r2}"].value
                        t2 = get_teacher(c, val2)
                        
                        count_t2 = t_counts.get(t2, 0)
                        cap_t2 = 2 if t2 == "IT_Dept" else 1
                        if t2 and count_t2 >= cap_t2: continue
                        
                        t_in_r2 = sum(1 for cx in all_classes if cx != c and is_academic(cx, r2) and get_teacher(cx, ws[f"{get_col(cx)}{r2}"].value) == t)
                        if t_in_r2 >= cap: continue
                        
                        ws[f"{get_col(c)}{r}"] = val2
                        ws[f"{get_col(c)}{r2}"] = val
                        t_counts[t] -= 1
                        if t2: t_counts[t2] = t_counts.get(t2, 0) + 1
                        break
    if not clash_found: break

for c in all_classes:
    for row in range(2, 67):
        wb[c][f"D{row}"] = ws[f"{get_col(c)}{row}"].value if ws[f"{get_col(c)}{row}"].value else ""
wb.save(file_path)
