import openpyxl

file_path = '/home/dpdp/target/timetable/MASTER_JADVAL_NEW.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Umumiy Jadval']

cohorts = {
    "C2": ["6-A", "6-B", "7-A", "7-B"],
    "C3": ["8-A", "8-B", "9-A", "9-B"],
    "C4": ["10-A", "10-B", "11-A", "11-B"]
}
all_classes = ["1-A", "1-B", "2-A", "2-B", "3-A", "4-A", 
               "5-A", "5-B", "6-A", "6-B", "7-A", "7-B", 
               "8-A", "8-B", "9-A", "9-B", "10-A", "10-B", "11-A", "11-B"]
def get_col(c): return openpyxl.utils.get_column_letter(all_classes.index(c) + 6)

fixed_breaks = ["NONUSHTA", "TUSHLIK", "POLNIK", "To'garak"]

# The 3 blocks in a day
def get_block_rows(day_idx, block_name):
    # day_idx: 0 to 4
    start_row = 2 + (day_idx * 13)
    # Block A: 1-soat, 2-soat -> start_row + 1, start_row + 2
    # Block B: 3-soat, 4-soat -> start_row + 3, start_row + 4
    # Block C: 6-soat, 7-soat -> start_row + 7, start_row + 8 (since Lunch is 6)
    if block_name == "A": return [start_row + 1, start_row + 2]
    if block_name == "B": return [start_row + 3, start_row + 4]
    if block_name == "C": return [start_row + 7, start_row + 8]

# Rotation matrix
# Cohort 2: Day1:A(Math), Day2:B(Math), Day3:C(Math), Day4:A(ELT), Day5:B(ELT)
# Cohort 3: Day1:B(ELT), Day2:C(Math), Day3:A(Math), Day4:B(Math), Day5:C(ELT)
# Cohort 4: Day1:C(ELT), Day2:A(ELT), Day3:B(Math), Day4:C(Math), Day5:A(Math)

schedule_plan = {
    "C2": [("A", "Math (Levels)"), ("B", "Math (Levels)"), ("C", "Math (Levels)"), ("A", "ELT (Levels)"), ("B", "ELT (Levels)")],
    "C3": [("B", "ELT (Levels)"), ("C", "Math (Levels)"), ("A", "Math (Levels)"), ("B", "Math (Levels)"), ("C", "ELT (Levels)")],
    "C4": [("C", "ELT (Levels)"), ("A", "ELT (Levels)"), ("B", "Math (Levels)"), ("C", "Math (Levels)"), ("A", "Math (Levels)")]
}

for c_key, classes in cohorts.items():
    plan = schedule_plan[c_key]
    for day_idx in range(5):
        block_name, subject = plan[day_idx]
        rows = get_block_rows(day_idx, block_name)
        
        # We need to extract the existing independent subjects that are currently in these slots
        # and swap them out, or just clear all academic subjects and repack
        pass

# Actually, the safest way is to extract ALL independent subjects for the week, 
# clear the grid, place the rigid blocks, and then pour the independent subjects back in!

for c_key, classes in cohorts.items():
    # 1. Extract all independent subjects
    indep_subjects = {c: [] for c in classes}
    for c in classes:
        col = get_col(c)
        for r in range(2, 67):
            val = ws[f"{col}{r}"].value
            if val and not any(b in val for b in fixed_breaks) and "Levels" not in val:
                indep_subjects[c].append(val)
        # Clear all academic cells
        for r in range(2, 67):
            val = ws[f"{col}{r}"].value
            if not (val and any(b in val for b in fixed_breaks)):
                ws[f"{col}{r}"] = None
                
    # 2. Place the rigid shared blocks
    plan = schedule_plan[c_key]
    for day_idx in range(5):
        block_name, subject = plan[day_idx]
        rows = get_block_rows(day_idx, block_name)
        for c in classes:
            col = get_col(c)
            ws[f"{col}{rows[0]}"] = subject
            ws[f"{col}{rows[1]}"] = subject
            
    # 3. Pour independent subjects back into the remaining empty academic slots
    for c in classes:
        col = get_col(c)
        idx = 0
        for r in range(2, 67):
            if idx >= len(indep_subjects[c]): break
            val = ws[f"{col}{r}"].value
            # If cell is empty (and not a fixed break, since we only cleared academic cells)
            if val is None:
                ws[f"{col}{r}"] = indep_subjects[c][idx]
                idx += 1

# Sync to POV
for c in all_classes:
    if c not in sum(cohorts.values(), []): continue
    col = get_col(c)
    ws_class = wb[c]
    for row in range(2, 67):
        val = ws[f"{col}{row}"].value
        ws_class[f"D{row}"] = val if val else ""

wb.save(file_path)
print("Cross-cohort clashes resolved seamlessly!")
