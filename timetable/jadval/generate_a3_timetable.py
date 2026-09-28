import os
import sys
import base64
import subprocess
import openpyxl
from collections import defaultdict

EXCEL_PATH = '/home/dpdp/Downloads/raspisanie_v2.xlsx'
LOGO_PATH = '/home/dpdp/target/assets/target-logo.png'
OUTPUT_DIR = '/home/dpdp/target/timetable/jadval'

# 20 classes
CLASSES = [
    '1A', '1B', '2A', '2B', '3A', '4A',
    '5A', '5B', '6A', '6B', '7A', '7B',
    '8A', '8B', '9A', '9B', '10A', '10B',
    '11A', '11B'
]

DAYS_CONFIG = [
    {'ru': 'Понедельник', 'uz': 'DUSHANBA', 'slug': '1_Dushanba'},
    {'ru': 'Вторник', 'uz': 'SESHANBA', 'slug': '2_Seshanba'},
    {'ru': 'Среда', 'uz': 'CHORSHANBA', 'slug': '3_Chorshanba'},
    {'ru': 'Четверг', 'uz': 'PAYSHANBA', 'slug': '4_Payshanba'},
    {'ru': 'Пятница', 'uz': 'JUMA', 'slug': '5_Juma'}
]

PERIOD_TIMES = {
    1: '09:00–09:40',
    2: '09:45–10:25',
    3: '10:30–11:10',
    4: '11:15–11:55',
    5: '12:00 / 12:45',
    6: '13:30–14:10',
    7: '14:15–14:55',
    8: '15:00–15:40',
    9: '16:10–16:50',
    10: '16:55–17:35'
}

SUBJECT_STYLES = {
    'Math': {'bg': '#eefcf1', 'border': '#bbf7d0', 'text': '#15803d', 'teacher': '#166534', 'short': 'Math'},
    'English': {'bg': '#fef2f2', 'border': '#fecaca', 'text': '#b91c1c', 'teacher': '#991b1b', 'short': 'Eng'},
    'IT': {'bg': '#eef2ff', 'border': '#c7d2fe', 'text': '#4338ca', 'teacher': '#3730a3', 'short': 'IT'},
    'Science': {'bg': '#f0f9ff', 'border': '#bae6fd', 'text': '#0369a1', 'teacher': '#075985', 'short': 'Sci'},
    'Physics': {'bg': '#f0f9ff', 'border': '#bae6fd', 'text': '#0369a1', 'teacher': '#075985', 'short': 'Phys'},
    'Chemistry': {'bg': '#f0f9ff', 'border': '#bae6fd', 'text': '#0369a1', 'teacher': '#075985', 'short': 'Chem'},
    'Biology': {'bg': '#f0f9ff', 'border': '#bae6fd', 'text': '#0369a1', 'teacher': '#075985', 'short': 'Bio'},
    'Uzbek': {'bg': '#fdf2f8', 'border': '#fbcfe8', 'text': '#be185d', 'teacher': '#9d174d', 'short': 'Uzb'},
    'Russian': {'bg': '#fdf2f8', 'border': '#fbcfe8', 'text': '#be185d', 'teacher': '#9d174d', 'short': 'Rus'},
    'Chinese': {'bg': '#fffbeb', 'border': '#fde68a', 'text': '#b45309', 'teacher': '#92400e', 'short': 'Chin'},
    'Global Perspective': {'bg': '#ecfeff', 'border': '#a5f3fc', 'text': '#0e7490', 'teacher': '#155e75', 'short': 'GP'},
    'G Pres': {'bg': '#ecfeff', 'border': '#a5f3fc', 'text': '#0e7490', 'teacher': '#155e75', 'short': 'GP'},
    'Economics': {'bg': '#f1f5f9', 'border': '#cbd5e1', 'text': '#334155', 'teacher': '#1e293b', 'short': 'Eco'},
    'Finance': {'bg': '#f1f5f9', 'border': '#cbd5e1', 'text': '#334155', 'teacher': '#1e293b', 'short': 'Fin'},
    'Business': {'bg': '#f1f5f9', 'border': '#cbd5e1', 'text': '#334155', 'teacher': '#1e293b', 'short': 'Bus'},
    'AI': {'bg': '#faf5ff', 'border': '#e9d5ff', 'text': '#7e22ce', 'teacher': '#6b21a8', 'short': 'AI'},
    'Per dev': {'bg': '#f0fdfa', 'border': '#99f6e4', 'text': '#0f766e', 'teacher': '#115e59', 'short': 'P.Dev'},
    'P.E / Gym': {'bg': '#f0fdf4', 'border': '#a7f3d0', 'text': '#047857', 'teacher': '#065f46', 'short': 'Gym'},
    'Chess': {'bg': '#fff7ed', 'border': '#fed7aa', 'text': '#c2410c', 'teacher': '#9a3412', 'short': 'Chess'},
    'Robotics': {'bg': '#fff7ed', 'border': '#fed7aa', 'text': '#c2410c', 'teacher': '#9a3412', 'short': 'Robot'},
    'Art': {'bg': '#fff1f2', 'border': '#fecdd3', 'text': '#9f1239', 'teacher': '#881337', 'short': 'Art'},
    'Music': {'bg': '#faf5ff', 'border': '#e9d5ff', 'text': '#86198f', 'teacher': '#701a75', 'short': 'Music'},
    'Mental.M': {'bg': '#faf5ff', 'border': '#e9d5ff', 'text': '#86198f', 'teacher': '#701a75', 'short': 'Mental'},
    'Homework': {'bg': '#f8fafc', 'border': '#e2e8f0', 'text': '#64748b', 'teacher': '#475569', 'short': 'HW'},
    'Choice': {'bg': '#f0fdf4', 'border': '#bbf7d0', 'text': '#166534', 'teacher': '#14532d', 'short': 'Choice'},
    'History': {'bg': '#fefce8', 'border': '#fef08a', 'text': '#854d0e', 'teacher': '#713f12', 'short': 'Hist'},
    'Hist.U': {'bg': '#fefce8', 'border': '#fef08a', 'text': '#854d0e', 'teacher': '#713f12', 'short': 'Hist.U'},
    'Hist.W': {'bg': '#fefce8', 'border': '#fef08a', 'text': '#854d0e', 'teacher': '#713f12', 'short': 'Hist.W'},
    'Geography': {'bg': '#fefce8', 'border': '#fef08a', 'text': '#854d0e', 'teacher': '#713f12', 'short': 'Geog'},
}

DEFAULT_STYLE = {'bg': '#f8fafc', 'border': '#e2e8f0', 'text': '#334155', 'teacher': '#64748b', 'short': '—'}

def get_style(subj_name):
    for key, style in SUBJECT_STYLES.items():
        if key.lower() in subj_name.lower():
            return style
    return DEFAULT_STYLE

def shorten_teacher(name):
    if not name:
        return ""
    name = name.strip()
    if 'vakant' in name.lower() or 'vacant' in name.lower():
        return 'Vakant'
    parts = name.split()
    if len(parts) >= 2:
        return f"{parts[0]} {parts[1][0]}."
    return name

def load_data():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb.active
    timetable = {d['ru']: {c: {p: [] for p in range(1, 11)} for c in CLASSES} for d in DAYS_CONFIG}
    for r in range(2, sheet.max_row + 1):
        day = str(sheet.cell(r, 1).value or '').strip()
        period = sheet.cell(r, 2).value
        cls_str = str(sheet.cell(r, 5).value or '').strip()
        subj = str(sheet.cell(r, 6).value or '').strip()
        teacher = str(sheet.cell(r, 7).value or '').strip()
        room = str(sheet.cell(r, 8).value or '').strip()
        if day not in timetable or not period:
            continue
        p = int(period)
        for c in [x.strip() for x in cls_str.split(',')]:
            if c in timetable[day]:
                timetable[day][c][p].append({
                    'subject': subj,
                    'teacher': teacher,
                    'room': room,
                    'cls_str': cls_str
                })
    return timetable

def format_cell_html(lessons):
    if not lessons:
        return '<div class="cell-empty">—</div>', DEFAULT_STYLE
    
    if len(lessons) == 1:
        l = lessons[0]
        subj = l['subject']
        teacher = shorten_teacher(l['teacher'])
        style = get_style(subj)
        disp_subj = subj
        if 'Global Perspective' in disp_subj:
            disp_subj = 'Global Persp.'
        elif 'Personal development' in disp_subj:
            disp_subj = 'Per. Dev.'
            
        html = f'''
        <div class="cell-inner">
            <div class="subj-name" style="color: {style['text']};">{disp_subj}</div>
            <div class="teacher-name" style="color: {style['teacher']};">{teacher}</div>
        </div>
        '''
        return html, style
    
    first_subj = lessons[0]['subject']
    base_subj = first_subj.split('·')[0].strip() if '·' in first_subj else first_subj
    style = get_style(base_subj)
    all_same_base = all(('·' in l['subject'] and l['subject'].split('·')[0].strip().lower() == base_subj.lower()) or (l['subject'].lower() == base_subj.lower()) for l in lessons)
    
    if all_same_base and base_subj.lower() in ['english', 'math']:
        teachers = [shorten_teacher(l['teacher']) for l in lessons]
        teachers_str = " · ".join(teachers)
        badge = f"{len(lessons)} gr."
        html = f'''
        <div class="cell-inner">
            <div class="subj-name" style="color: {style['text']};">{base_subj} <span class="badge" style="background:{style['border']}; color:{style['text']};">{badge}</span></div>
            <div class="teacher-name" style="color: {style['teacher']};">{teachers_str}</div>
        </div>
        '''
        return html, style
    elif 'choice' in first_subj.lower():
        style = get_style('Choice')
        html = f'''
        <div class="cell-inner">
            <div class="subj-name" style="color: {style['text']};">Choice / Tanlov <span class="badge" style="background:#bbf7d0; color:#15803d;">5 yo'n.</span></div>
            <div class="teacher-name" style="color: {style['teacher']};">Phys · Chem · IT · Math · Eng</div>
        </div>
        '''
        return html, style
    elif 'it' in first_subj.lower():
        style = get_style('IT')
        html = f'''
        <div class="cell-inner">
            <div class="subj-name" style="color: {style['text']};">IT Streams <span class="badge" style="background:#c7d2fe; color:#4338ca;">{len(lessons)} gr.</span></div>
            <div class="teacher-name" style="color: {style['teacher']};">Cyber · Vibe · Mobile · Python</div>
        </div>
        '''
        return html, style
    else:
        subjs_str = " / ".join(set(l['subject'].split('·')[0].strip() for l in lessons))
        teachers = [shorten_teacher(l['teacher']) for l in lessons]
        html = f'''
        <div class="cell-inner">
            <div class="subj-name" style="color: {style['text']};">{subjs_str}</div>
            <div class="teacher-name" style="color: {style['teacher']};">{" · ".join(teachers)}</div>
        </div>
        '''
        return html, style

def get_base64_logo():
    if os.path.exists(LOGO_PATH):
        with open(LOGO_PATH, 'rb') as f:
            return f"data:image/png;base64,{base64.b64encode(f.read()).decode('utf-8')}"
    return ""

def generate_day_page_html(day_info, timetable, logo_b64):
    day_ru = day_info['ru']
    day_uz = day_info['uz']
    
    rows_html = []
    for cls in CLASSES:
        is_primary = cls in ['1A', '1B', '2A', '2B', '3A', '4A']
        cells_html = []
        for p in range(1, 11):
            lessons = timetable[day_ru][cls][p]
            cell_content, style = format_cell_html(lessons)
            extra_class = " period-5" if p == 5 else ""
            cells_html.append(f'''
            <td class="cell{extra_class}" style="background-color: {style['bg']}; border-color: {style['border']};">
                {cell_content}
            </td>
            ''')
            
        cls_badge_class = "cls-primary" if is_primary else "cls-secondary"
        row_class = "row-grade-break" if cls in ['4A', '8B', '11B'] else ""
        
        rows_html.append(f'''
        <tr class="{row_class}">
            <th class="cell-cls {cls_badge_class}">
                <div class="cls-label">{cls}</div>
            </th>
            {"".join(cells_html)}
        </tr>
        ''')
        
    cols_header = []
    for p in range(1, 11):
        time_str = PERIOD_TIMES[p]
        if p == 5:
            time_str = '<span title="1-4 Lunch / 5-11 Lesson">12:00</span> / <span title="5-11 Lunch / 1-4 Lesson">12:45</span>'
        cols_header.append(f'''
        <th class="col-period">
            <div class="p-num">{p}-DARS</div>
            <div class="p-time">{time_str}</div>
        </th>
        ''')
        
    page_html = f'''
    <div class="page">
        <!-- HEADER -->
        <div class="header">
            <div class="header-left">
                {f'<img src="{logo_b64}" class="logo" alt="Target Logo" />' if logo_b64 else ''}
                <div class="school-info">
                    <div class="school-name">TARGET INTERNATIONAL SCHOOL</div>
                    <div class="branch-name">Yunusobod Filiali · 2026–2027 O'quv Yili</div>
                </div>
            </div>
            
            <div class="header-center">
                <div class="timetable-title">UMUMIY DARS JADVALI</div>
                <div class="schedule-pills">
                    <span class="pill pill-lunch">1–4 Tushlik: 12:00–12:40</span>
                    <span class="pill pill-lunch">5–11 Tushlik: 12:45–13:25</span>
                    <span class="pill pill-snack">Poldnik: 15:40–16:10</span>
                </div>
            </div>
            
            <div class="header-right">
                <div class="day-badge">
                    <span class="day-uz">{day_uz}</span>
                    <span class="day-ru">{day_ru}</span>
                </div>
            </div>
        </div>

        <!-- MAIN TABLE -->
        <div class="table-container">
            <table class="timetable-table">
                <thead>
                    <tr>
                        <th class="col-cls">
                            <div class="p-num">SINF</div>
                            <div class="p-time">КЛАСС</div>
                        </th>
                        {"".join(cols_header)}
                    </tr>
                </thead>
                <tbody>
                    {"".join(rows_html)}
                </tbody>
            </table>
        </div>

        <!-- FOOTER / LEGEND -->
        <div class="footer">
            <div class="legend">
                <span class="leg-item"><span class="leg-dot" style="background:#bbf7d0;"></span> Math</span>
                <span class="leg-item"><span class="leg-dot" style="background:#fecaca;"></span> English</span>
                <span class="leg-item"><span class="leg-dot" style="background:#c7d2fe;"></span> IT Streams</span>
                <span class="leg-item"><span class="leg-dot" style="background:#bae6fd;"></span> Science / Physics</span>
                <span class="leg-item"><span class="leg-dot" style="background:#fbcfe8;"></span> Uzbek / Russian</span>
                <span class="leg-item"><span class="leg-dot" style="background:#fde68a;"></span> Chinese / History</span>
                <span class="leg-item"><span class="leg-dot" style="background:#a5f3fc;"></span> Global Persp.</span>
                <span class="leg-item"><span class="leg-dot" style="background:#cbd5e1;"></span> Eco / Finance</span>
                <span class="leg-item"><span class="leg-dot" style="background:#e9d5ff;"></span> AI / Art / Music</span>
                <span class="leg-item"><span class="leg-dot" style="background:#a7f3d0;"></span> P.E / Gym</span>
                <span class="leg-item"><span class="leg-dot" style="background:#fed7aa;"></span> Robotics / Chess</span>
            </div>
            <div class="page-meta">
                Format: A3 Landscape (420 × 297 mm) · Chop etish uchun tayyor (A3 Ready)
            </div>
        </div>
    </div>
    '''
    return page_html

def build_full_html(pages_html):
    return f'''<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <title>Dars Jadvali A3</title>
    <style>
        @page {{
            size: 420mm 297mm;
            margin: 6mm 8mm;
        }}
        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            background: #ffffff;
            color: #0f172a;
            -webkit-font-smoothing: antialiased;
        }}
        .page {{
            width: 404mm;
            height: 284mm;
            max-height: 284mm;
            display: flex;
            flex-direction: column;
            justify-content: space-between;
            page-break-after: always;
            break-after: page;
            overflow: hidden;
        }}
        .page:last-child {{
            page-break-after: auto;
            break-after: auto;
        }}
        .header {{
            height: 20mm;
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-bottom: 2px solid #0f172a;
            padding-bottom: 2mm;
            margin-bottom: 2mm;
        }}
        .header-left {{
            display: flex;
            align-items: center;
            gap: 10px;
        }}
        .logo {{
            height: 16mm;
            width: auto;
            object-fit: contain;
        }}
        .school-info {{
            display: flex;
            flex-direction: column;
        }}
        .school-name {{
            font-size: 16px;
            font-weight: 900;
            letter-spacing: 0.5px;
            color: #0f172a;
            text-transform: uppercase;
        }}
        .branch-name {{
            font-size: 11px;
            font-weight: 600;
            color: #64748b;
        }}
        .header-center {{
            display: flex;
            flex-direction: column;
            align-items: center;
            gap: 3px;
        }}
        .timetable-title {{
            font-size: 15px;
            font-weight: 800;
            color: #0f172a;
            letter-spacing: 1px;
        }}
        .schedule-pills {{
            display: flex;
            gap: 8px;
        }}
        .pill {{
            font-size: 9px;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 12px;
            border: 1px solid transparent;
        }}
        .pill-lunch {{
            background: #fef3c7;
            color: #92400e;
            border-color: #fde68a;
        }}
        .pill-snack {{
            background: #e0f2fe;
            color: #0369a1;
            border-color: #bae6fd;
        }}
        .header-right {{
            display: flex;
            align-items: center;
        }}
        .day-badge {{
            background: #0f172a;
            color: #ffffff;
            padding: 4px 18px;
            border-radius: 6px;
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
        }}
        .day-uz {{
            font-size: 15px;
            font-weight: 900;
            letter-spacing: 1px;
            line-height: 1.1;
        }}
        .day-ru {{
            font-size: 9px;
            font-weight: 600;
            color: #94a3b8;
            letter-spacing: 0.5px;
            text-transform: uppercase;
        }}
        .table-container {{
            flex: 1;
            display: flex;
            width: 100%;
        }}
        .timetable-table {{
            width: 100%;
            height: 100%;
            border-collapse: collapse;
            table-layout: fixed;
            border: 1.5px solid #0f172a;
        }}
        .timetable-table th,
        .timetable-table td {{
            border: 1px solid #cbd5e1;
            padding: 0;
            text-align: center;
            vertical-align: middle;
        }}
        .timetable-table thead tr {{
            height: 10mm;
            background: #0f172a;
            color: #ffffff;
        }}
        .col-cls {{
            width: 18mm;
            background: #0f172a;
            color: #ffffff;
        }}
        .col-period {{
            width: 38.6mm;
            background: #0f172a;
            color: #ffffff;
            border-right: 1px solid #334155 !important;
        }}
        .p-num {{
            font-size: 11px;
            font-weight: 800;
            letter-spacing: 0.5px;
            line-height: 1.1;
        }}
        .p-time {{
            font-size: 8px;
            font-weight: 500;
            color: #94a3b8;
            margin-top: 1px;
            line-height: 1;
        }}
        .timetable-table tbody tr {{
            height: 12.3mm;
        }}
        .row-grade-break td,
        .row-grade-break th {{
            border-bottom: 2px solid #64748b !important;
        }}
        .cell-cls {{
            background: #f1f5f9;
            border-right: 2px solid #0f172a !important;
        }}
        .cls-primary {{
            background: #eff6ff;
        }}
        .cls-secondary {{
            background: #f8fafc;
        }}
        .cls-label {{
            font-size: 13px;
            font-weight: 900;
            color: #0f172a;
            letter-spacing: 0.5px;
        }}
        .cell {{
            padding: 2px 4px !important;
        }}
        .cell-inner {{
            width: 100%;
            height: 100%;
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            overflow: hidden;
        }}
        .subj-name {{
            font-size: 9.5px;
            font-weight: 800;
            line-height: 1.15;
            white-space: nowrap;
            overflow: hidden;
            text-overflow: ellipsis;
            max-width: 100%;
            display: flex;
            align-items: center;
            gap: 3px;
        }}
        .teacher-name {{
            font-size: 7.5px;
            font-weight: 600;
            line-height: 1.1;
            margin-top: 1.5px;
            white-space: nowrap;
            overflow: hidden;
            text-overflow: ellipsis;
            max-width: 100%;
        }}
        .badge {{
            font-size: 6.5px;
            font-weight: 800;
            padding: 0.5px 3px;
            border-radius: 4px;
            text-transform: uppercase;
        }}
        .cell-empty {{
            font-size: 11px;
            color: #cbd5e1;
            font-weight: bold;
        }}
        .footer {{
            height: 6mm;
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-top: 1px solid #cbd5e1;
            margin-top: 1.5mm;
            padding-top: 1mm;
        }}
        .legend {{
            display: flex;
            align-items: center;
            gap: 10px;
        }}
        .leg-item {{
            display: flex;
            align-items: center;
            gap: 3px;
            font-size: 8px;
            font-weight: 600;
            color: #475569;
        }}
        .leg-dot {{
            width: 8px;
            height: 8px;
            border-radius: 2px;
            border: 1px solid rgba(0,0,0,0.1);
            display: inline-block;
        }}
        .page-meta {{
            font-size: 8px;
            font-weight: 600;
            color: #94a3b8;
        }}
    </style>
</head>
<body>
    {''.join(pages_html)}
</body>
</html>
'''

def generate_full_week_page_html(timetable, logo_b64):
    """50 periods across 5 days (Mon-Fri) on a single A3 master sheet, removing Saturday & empty margins."""
    rows_html = []
    for cls in CLASSES:
        is_primary = cls in ['1A', '1B', '2A', '2B', '3A', '4A']
        cells_html = []
        for d in DAYS_CONFIG:
            day_ru = d['ru']
            for p in range(1, 11):
                lessons = timetable[day_ru][cls][p]
                if not lessons:
                    cells_html.append('<td class="w-cell cell-empty">—</td>')
                    continue
                first_subj = lessons[0]['subject']
                base_subj = first_subj.split('·')[0].strip() if '·' in first_subj else first_subj
                style = get_style(base_subj)
                
                # Abbreviated subject
                short_s = style.get('short', base_subj[:4])
                teacher_s = shorten_teacher(lessons[0]['teacher'])
                
                # Multi-group indicator
                grp_tag = f" <sup>+{len(lessons)-1}</sup>" if len(lessons) > 1 else ""
                
                # Last period of day gets a divider border
                border_r = " border-day-end" if p == 10 else ""
                cells_html.append(f'''
                <td class="w-cell{border_r}" style="background-color: {style['bg']}; color: {style['text']};">
                    <div class="w-subj">{short_s}{grp_tag}</div>
                    <div class="w-teach" style="color: {style['teacher']};">{teacher_s}</div>
                </td>
                ''')
                
        cls_badge_class = "cls-primary" if is_primary else "cls-secondary"
        row_class = "row-grade-break" if cls in ['4A', '8B', '11B'] else ""
        rows_html.append(f'''
        <tr class="{row_class}">
            <th class="w-col-cls {cls_badge_class}">
                <div class="cls-label">{cls}</div>
            </th>
            {"".join(cells_html)}
        </tr>
        ''')
        
    day_headers = []
    sub_headers = []
    for d in DAYS_CONFIG:
        day_headers.append(f'''
        <th colspan="10" class="w-day-header">
            {d['uz']} / {d['ru']}
        </th>
        ''')
        for p in range(1, 11):
            sub_headers.append(f'''
            <th class="w-p-header{' border-day-end' if p == 10 else ''}">{p}</th>
            ''')
            
    html = f'''<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <title>Haftalik Dars Jadvali (Full Week A3)</title>
    <style>
        @page {{
            size: 420mm 297mm;
            margin: 5mm 6mm;
        }}
        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            background: #ffffff;
            color: #0f172a;
        }}
        .page {{
            width: 408mm;
            height: 287mm;
            display: flex;
            flex-direction: column;
            justify-content: space-between;
            overflow: hidden;
        }}
        .header {{
            height: 16mm;
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-bottom: 2px solid #0f172a;
            padding-bottom: 1.5mm;
            margin-bottom: 1.5mm;
        }}
        .header-left {{
            display: flex;
            align-items: center;
            gap: 8px;
        }}
        .logo {{
            height: 13mm;
            width: auto;
        }}
        .school-name {{
            font-size: 14px;
            font-weight: 900;
            letter-spacing: 0.5px;
        }}
        .branch-name {{
            font-size: 10px;
            font-weight: 600;
            color: #64748b;
        }}
        .title {{
            font-size: 14px;
            font-weight: 800;
            letter-spacing: 1px;
        }}
        .pills {{
            display: flex;
            gap: 6px;
        }}
        .pill {{
            font-size: 8px;
            font-weight: 600;
            padding: 1.5px 6px;
            border-radius: 10px;
            background: #f1f5f9;
        }}
        .table-container {{
            flex: 1;
            display: flex;
            width: 100%;
        }}
        .w-table {{
            width: 100%;
            height: 100%;
            border-collapse: collapse;
            table-layout: fixed;
            border: 1.5px solid #0f172a;
        }}
        .w-table th, .w-table td {{
            border: 0.5px solid #cbd5e1;
            padding: 0;
            text-align: center;
            vertical-align: middle;
        }}
        .w-col-cls {{
            width: 14mm;
            background: #0f172a;
            color: #ffffff;
            border-right: 1.5px solid #0f172a !important;
        }}
        .cls-label {{
            font-size: 11px;
            font-weight: 800;
            color: #0f172a;
        }}
        .w-day-header {{
            background: #0f172a;
            color: #ffffff;
            font-size: 10px;
            font-weight: 800;
            height: 5.5mm;
            border-right: 1.5px solid #ffffff !important;
        }}
        .w-p-header {{
            background: #1e293b;
            color: #cbd5e1;
            font-size: 8px;
            font-weight: 700;
            height: 4mm;
        }}
        .border-day-end {{
            border-right: 1.5px solid #0f172a !important;
        }}
        .w-table tbody tr {{
            height: 12.8mm;
        }}
        .row-grade-break td, .row-grade-break th {{
            border-bottom: 1.5px solid #475569 !important;
        }}
        .w-cell {{
            padding: 1px !important;
            font-size: 7.5px;
            overflow: hidden;
        }}
        .w-subj {{
            font-weight: 800;
            line-height: 1.1;
            white-space: nowrap;
            overflow: hidden;
            text-overflow: ellipsis;
        }}
        .w-teach {{
            font-size: 6px;
            font-weight: 600;
            line-height: 1;
            margin-top: 1px;
            white-space: nowrap;
            overflow: hidden;
            text-overflow: ellipsis;
        }}
        .footer {{
            height: 5mm;
            display: flex;
            align-items: center;
            justify-content: space-between;
            border-top: 1px solid #cbd5e1;
            margin-top: 1mm;
            font-size: 7.5px;
            color: #64748b;
        }}
    </style>
</head>
<body>
    <div class="page">
        <div class="header">
            <div class="header-left">
                {f'<img src="{logo_b64}" class="logo" />' if logo_b64 else ''}
                <div>
                    <div class="school-name">TARGET INTERNATIONAL SCHOOL</div>
                    <div class="branch-name">Yunusobod Filiali · 2026–2027 O'quv Yili</div>
                </div>
            </div>
            <div class="title">HAFTALIK DARS JADVALI (MASTER SCHEDULE)</div>
            <div class="pills">
                <span class="pill">1–4 Tushlik: 12:00–12:40</span>
                <span class="pill">5–11 Tushlik: 12:45–13:25</span>
                <span class="pill">Poldnik: 15:40–16:10</span>
            </div>
        </div>
        <div class="table-container">
            <table class="w-table">
                <thead>
                    <tr>
                        <th rowspan="2" class="w-col-cls" style="color:white; font-size:9px;">SINF</th>
                        {"".join(day_headers)}
                    </tr>
                    <tr>
                        {"".join(sub_headers)}
                    </tr>
                </thead>
                <tbody>
                    {"".join(rows_html)}
                </tbody>
            </table>
        </div>
        <div class="footer">
            <div>5 kunlik to'liq jadval (Dushanba–Juma) · Shanba darslari yo'q (Removed)</div>
            <div>A3 Master Format (420 × 297 mm) · Zero Unnecessary Space</div>
        </div>
    </div>
</body>
</html>
'''
    return html

def main():
    print("Loading timetable data from Excel...")
    timetable = load_data()
    logo_b64 = get_base64_logo()
    
    # 1. Generate All 5 Days Master Document (5 Pages A3)
    print("Generating 5-Day A3 Master PDF (1 Day Per Page)...")
    all_pages_html = [generate_day_page_html(d, timetable, logo_b64) for d in DAYS_CONFIG]
    all_full_html = build_full_html(all_pages_html)
    all_html_file = os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_AllDays.html')
    with open(all_html_file, 'w', encoding='utf-8') as f:
        f.write(all_full_html)
        
    all_pdf_file = os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_AllDays.pdf')
    subprocess.run([
        'google-chrome', '--headless', '--disable-gpu',
        f'--print-to-pdf={all_pdf_file}', '--print-to-pdf-no-header',
        all_html_file
    ], check=True)
    print(f" -> Generated: {all_pdf_file}")
    
    # 2. Generate Individual Days (Monday to Friday)
    for d in DAYS_CONFIG:
        slug = d['slug']
        page_html = generate_day_page_html(d, timetable, logo_b64)
        single_full_html = build_full_html([page_html])
        
        single_html_file = os.path.join(OUTPUT_DIR, f'raspisanie_v2_matrix_A3_{slug}.html')
        with open(single_html_file, 'w', encoding='utf-8') as f:
            f.write(single_full_html)
            
        single_pdf_file = os.path.join(OUTPUT_DIR, f'raspisanie_v2_matrix_A3_{slug}.pdf')
        subprocess.run([
            'google-chrome', '--headless', '--disable-gpu',
            f'--print-to-pdf={single_pdf_file}', '--print-to-pdf-no-header',
            single_html_file
        ], check=True)
        print(f" -> Generated: {single_pdf_file}")

    # Copy Monday to raspisanie_v2_matrix_A3_Monday.pdf and raspisanie_v2_matrix_A3.pdf
    subprocess.run(['cp', os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_1_Dushanba.pdf'), os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_Monday.pdf')], check=True)
    subprocess.run(['cp', os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_1_Dushanba.pdf'), os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3.pdf')], check=True)

    # 3. Generate Full Week 5-Day Matrix (Single A3 Landscape Sheet)
    print("Generating Full Week Master A3 (Single Sheet)...")
    full_week_html = generate_full_week_page_html(timetable, logo_b64)
    fw_html_file = os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_FullWeek.html')
    with open(fw_html_file, 'w', encoding='utf-8') as f:
        f.write(full_week_html)
        
    fw_pdf_file = os.path.join(OUTPUT_DIR, 'raspisanie_v2_matrix_A3_FullWeek.pdf')
    subprocess.run([
        'google-chrome', '--headless', '--disable-gpu',
        f'--print-to-pdf={fw_pdf_file}', '--print-to-pdf-no-header',
        fw_html_file
    ], check=True)
    print(f" -> Generated: {fw_pdf_file}")
    
    print("\nAll A3 Timetables successfully generated!")

if __name__ == '__main__':
    main()
