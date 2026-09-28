import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_gradebook_excel(filename="Teacher_Production_Gradebook.xlsx"):
    wb = openpyxl.Workbook()

    font_family = "Arial"
    header_title_font = Font(name=font_family, size=15, bold=True, color="0F172A")
    sub_title_font = Font(name=font_family, size=9, italic=True, color="64748B")
    
    month_header_font = Font(name=font_family, size=11, bold=True, color="0F172A")
    lesson_header_font = Font(name=font_family, size=9, bold=True, color="1E1B4B")
    cat_header_font = Font(name=font_family, size=8, bold=True, color="334155")
    
    regular_font = Font(name=font_family, size=9, color="0F172A")
    bold_font = Font(name=font_family, size=9, bold=True, color="0F172A")
    formula_font = Font(name=font_family, size=9, bold=True, color="1E40AF")
    
    month_fill_even = PatternFill(start_color="EEF2F6", end_color="EEF2F6", fill_type="solid")
    month_fill_odd = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    lesson_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    day_total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    month_total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    month_pct_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    overall_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    thick_right = Border(
        right=Side(style='medium', color='64748B'),
        left=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws = wb.active
    ws.title = "Gradebook"
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws['A1'] = "CLASS GRADEBOOK & LESSON SCHEDULE JOURNAL"
    ws['A1'].font = header_title_font
    ws['A2'] = "Per-Lesson Evaluations: Categories evaluated on each lesson day | Academic Year: 2026-2027"
    ws['A2'].font = sub_title_font

    categories = [
        ("Attendance", 5),
        ("Participation", 10),
        ("Homework", 10)
    ]
    single_lesson_cap = sum(c[1] for c in categories)

    months_schedule = [
        ("September", [("Day 1", "Sep 02"), ("Day 2", "Sep 05"), ("Day 3", "Sep 09"), ("Day 4", "Sep 12")]),
        ("October",   [("Day 5", "Oct 03"), ("Day 6", "Oct 07"), ("Day 7", "Oct 10"), ("Day 8", "Oct 14")]),
        ("November",  [("Day 9", "Nov 04"), ("Day 10", "Nov 07"), ("Day 11", "Nov 11"), ("Day 12", "Nov 14")]),
    ]

    # Header Row 3: Roll & Name
    ws['A3'] = "Roll"
    ws['A4'] = "ID"
    ws['A5'] = "#"
    ws['B3'] = "Student Name"
    ws['B4'] = "Full Name"
    ws['B5'] = "Name"

    for r in [3, 4, 5]:
        for c in ['A', 'B']:
            ws[f'{c}{r}'].font = bold_font
            ws[f'{c}{r}'].alignment = Alignment(horizontal='center', vertical='center')
            ws[f'{c}{r}'].fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            ws[f'{c}{r}'].border = thin_border

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24

    col_idx = 3
    month_group_ranges = []
    month_pct_cols = []

    for m_idx, (month_name, lessons) in enumerate(months_schedule):
        m_start_col = col_idx
        m_fill = month_fill_odd if m_idx % 2 == 0 else month_fill_even
        month_cap = len(lessons) * single_lesson_cap

        for lesson_title, lesson_date in lessons:
            l_start_col = col_idx

            # Categories for this lesson
            for cat_name, cap in categories:
                c_let = get_column_letter(col_idx)
                ws[f'{c_let}5'] = f"{cat_name}\n(/{cap})"
                ws[f'{c_let}5'].font = cat_header_font
                ws[f'{c_let}5'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws[f'{c_let}5'].fill = m_fill
                ws[f'{c_let}5'].border = thin_border
                ws.column_dimensions[c_let].width = 11
                col_idx += 1

            # Day Total
            dt_let = get_column_letter(col_idx)
            ws[f'{dt_let}5'] = f"Total\n(/{single_lesson_cap})"
            ws[f'{dt_let}5'].font = bold_font
            ws[f'{dt_let}5'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws[f'{dt_let}5'].fill = day_total_fill
            ws[f'{dt_let}5'].border = thick_right
            ws.column_dimensions[dt_let].width = 9
            col_idx += 1

            l_end_col = col_idx - 1

            # Merge Lesson Header Row 4
            ls_let = get_column_letter(l_start_col)
            le_let = get_column_letter(l_end_col)
            ws.merge_cells(f'{ls_let}4:{le_let}4')
            ws[f'{ls_let}4'] = f"{lesson_title} ({lesson_date})"
            ws[f'{ls_let}4'].font = lesson_header_font
            ws[f'{ls_let}4'].alignment = Alignment(horizontal='center', vertical='center')
            ws[f'{ls_let}4'].fill = lesson_fill
            for c in range(l_start_col, l_end_col + 1):
                ws.cell(row=4, column=c).border = thin_border
                ws.cell(row=4, column=c).fill = lesson_fill
            ws.cell(row=4, column=l_end_col).border = thick_right

        # Month Total Column
        mt_let = get_column_letter(col_idx)
        ws[f'{mt_let}4'] = f"{month_name[:3]} Total"
        ws[f'{mt_let}5'] = f"(/{month_cap})"
        ws[f'{mt_let}4'].font = bold_font
        ws[f'{mt_let}5'].font = bold_font
        ws[f'{mt_let}4'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{mt_let}5'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{mt_let}4'].fill = month_total_fill
        ws[f'{mt_let}5'].fill = month_total_fill
        ws[f'{mt_let}4'].border = thin_border
        ws[f'{mt_let}5'].border = thin_border
        ws.column_dimensions[mt_let].width = 11
        col_idx += 1

        # Month % Column
        mp_let = get_column_letter(col_idx)
        month_pct_cols.append(mp_let)
        ws[f'{mp_let}4'] = f"{month_name[:3]} %"
        ws[f'{mp_let}5'] = "Score %"
        ws[f'{mp_let}4'].font = bold_font
        ws[f'{mp_let}5'].font = bold_font
        ws[f'{mp_let}4'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{mp_let}5'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{mp_let}4'].fill = month_pct_fill
        ws[f'{mp_let}5'].fill = month_pct_fill
        ws[f'{mp_let}4'].border = thick_right
        ws[f'{mp_let}5'].border = thick_right
        ws.column_dimensions[mp_let].width = 12
        col_idx += 1

        m_end_col = col_idx - 1

        # Merge Month Header Row 3
        ms_let = get_column_letter(m_start_col)
        me_let = get_column_letter(m_end_col)
        ws.merge_cells(f'{ms_let}3:{me_let}3')
        ws[f'{ms_let}3'] = f"📅 {month_name} ({len(lessons)} Lessons)"
        ws[f'{ms_let}3'].font = month_header_font
        ws[f'{ms_let}3'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'{ms_let}3'].fill = m_fill
        for c in range(m_start_col, m_end_col + 1):
            ws.cell(row=3, column=c).border = thin_border
            ws.cell(row=3, column=c).fill = m_fill
        ws.cell(row=3, column=m_end_col).border = thick_right

        # Save range for expandable grouping (group all lessons so month collapses to Month Total & %)
        cat_group_end = col_idx - 3
        month_group_ranges.append((get_column_letter(m_start_col), get_column_letter(cat_group_end)))

    # Overall Columns
    ov_avg_let = get_column_letter(col_idx)
    ws[f'{ov_avg_let}3'] = "OVERALL"
    ws[f'{ov_avg_let}4'] = "Average %"
    ws[f'{ov_avg_let}5'] = "YTD"
    ws[f'{ov_avg_let}3'].font = bold_font
    ws[f'{ov_avg_let}4'].font = bold_font
    ws[f'{ov_avg_let}5'].font = bold_font
    ws[f'{ov_avg_let}3'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{ov_avg_let}4'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{ov_avg_let}5'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{ov_avg_let}3'].fill = overall_fill
    ws[f'{ov_avg_let}4'].fill = overall_fill
    ws[f'{ov_avg_let}5'].fill = overall_fill
    ws[f'{ov_avg_let}3'].border = thin_border
    ws[f'{ov_avg_let}4'].border = thin_border
    ws[f'{ov_avg_let}5'].border = thin_border
    ws.column_dimensions[ov_avg_let].width = 15
    col_idx += 1

    tier_let = get_column_letter(col_idx)
    ws[f'{tier_let}3'] = "STANDING"
    ws[f'{tier_let}4'] = "Grade Tier"
    ws[f'{tier_let}5'] = "Tier"
    ws[f'{tier_let}3'].font = bold_font
    ws[f'{tier_let}4'].font = bold_font
    ws[f'{tier_let}5'].font = bold_font
    ws[f'{tier_let}3'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{tier_let}4'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{tier_let}5'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'{tier_let}3'].fill = overall_fill
    ws[f'{tier_let}4'].fill = overall_fill
    ws[f'{tier_let}5'].fill = overall_fill
    ws[f'{tier_let}3'].border = thin_border
    ws[f'{tier_let}4'].border = thin_border
    ws[f'{tier_let}5'].border = thin_border
    ws.column_dimensions[tier_let].width = 12

    # 35 Student Rows
    start_row = 6
    num_students = 35

    for r in range(start_row, start_row + num_students):
        s_num = r - start_row + 1
        ws[f'A{r}'] = s_num
        ws[f'A{r}'].alignment = Alignment(horizontal='center')
        ws[f'A{r}'].font = regular_font
        ws[f'A{r}'].border = thin_border

        ws[f'B{r}'].font = bold_font
        ws[f'B{r}'].border = thin_border

        c_idx = 3
        for month_name, lessons in months_schedule:
            month_cap = len(lessons) * single_lesson_cap
            day_total_cells = []

            for _ in lessons:
                first_cat = get_column_letter(c_idx)
                last_cat = get_column_letter(c_idx + len(categories) - 1)

                for _ in categories:
                    cell = ws.cell(row=r, column=c_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='center')
                    cell.font = regular_font
                    c_idx += 1

                # Day Total Formula
                dt_cell = ws.cell(row=r, column=c_idx)
                dt_cell.value = f"=SUM({first_cat}{r}:{last_cat}{r})"
                dt_cell.font = formula_font
                dt_cell.fill = day_total_fill
                dt_cell.border = thick_right
                dt_cell.alignment = Alignment(horizontal='center')
                day_total_cells.append(f"{get_column_letter(c_idx)}{r}")
                c_idx += 1

            # Month Total Formula
            mt_cell = ws.cell(row=r, column=c_idx)
            mt_let = get_column_letter(c_idx)
            mt_cell.value = f"=SUM({','.join(day_total_cells)})"
            mt_cell.font = formula_font
            mt_cell.fill = month_total_fill
            mt_cell.border = thin_border
            mt_cell.alignment = Alignment(horizontal='center')
            c_idx += 1

            # Month % Formula
            mp_cell = ws.cell(row=r, column=c_idx)
            mp_cell.value = f"=IF({mt_let}{r}>0, {mt_let}{r}/{month_cap}, 0)"
            mp_cell.number_format = '0.0%'
            mp_cell.font = formula_font
            mp_cell.fill = month_pct_fill
            mp_cell.border = thick_right
            mp_cell.alignment = Alignment(horizontal='center')
            c_idx += 1

        # Overall Average Formula
        pct_refs = [f"{col_let}{r}" for col_let in month_pct_cols]
        ov_cell = ws.cell(row=r, column=c_idx)
        ov_cell.value = f"=AVERAGE({','.join(pct_refs)})"
        ov_cell.number_format = '0.0%'
        ov_cell.font = bold_font
        ov_cell.fill = overall_fill
        ov_cell.border = thin_border
        ov_cell.alignment = Alignment(horizontal='center')

        # Standing Formula
        t_cell = ws.cell(row=r, column=c_idx + 1)
        t_cell.value = f'=IF({ov_avg_let}{r}>=0.9,"A",IF({ov_avg_let}{r}>=0.8,"B",IF({ov_avg_let}{r}>=0.7,"C",IF({ov_avg_let}{r}>=0.6,"D","F")))'
        t_cell.font = bold_font
        t_cell.fill = overall_fill
        t_cell.border = thin_border
        t_cell.alignment = Alignment(horizontal='center')

    # Apply Native Column Grouping for Google Sheets [+] and [-] expand/collapse
    for g_start, g_end in month_group_ranges:
        ws.column_dimensions.group(g_start, g_end, hidden=False)

    # Class Average Row
    avg_row = start_row + num_students
    ws[f'A{avg_row}'] = "—"
    ws[f'B{avg_row}'] = "CLASS AVERAGE"
    ws[f'B{avg_row}'].font = Font(name=font_family, size=10, bold=True, color="1E3A8A")
    ws[f'A{avg_row}'].border = thin_border
    ws[f'B{avg_row}'].border = thin_border

    c_idx = 3
    for month_name, lessons in months_schedule:
        for _ in lessons:
            for _ in categories:
                c_let = get_column_letter(c_idx)
                c = ws.cell(row=avg_row, column=c_idx)
                c.value = f"=AVERAGE({c_let}{start_row}:{c_let}{avg_row-1})"
                c.font = bold_font
                c.border = thin_border
                c.alignment = Alignment(horizontal='center')
                c_idx += 1
            # Day total avg
            c_let = get_column_letter(c_idx)
            c = ws.cell(row=avg_row, column=c_idx)
            c.value = f"=AVERAGE({c_let}{start_row}:{c_let}{avg_row-1})"
            c.font = bold_font
            c.fill = day_total_fill
            c.border = thick_right
            c.alignment = Alignment(horizontal='center')
            c_idx += 1

        # Month total avg
        c_let = get_column_letter(c_idx)
        c = ws.cell(row=avg_row, column=c_idx)
        c.value = f"=AVERAGE({c_let}{start_row}:{c_let}{avg_row-1})"
        c.font = bold_font
        c.fill = month_total_fill
        c.border = thin_border
        c.alignment = Alignment(horizontal='center')
        c_idx += 1

        # Month % avg
        c_let = get_column_letter(c_idx)
        c = ws.cell(row=avg_row, column=c_idx)
        c.value = f"=AVERAGE({c_let}{start_row}:{c_let}{avg_row-1})"
        c.number_format = '0.0%'
        c.font = bold_font
        c.fill = month_pct_fill
        c.border = thick_right
        c.alignment = Alignment(horizontal='center')
        c_idx += 1

    # Overall Avg of Class
    c_let = get_column_letter(c_idx)
    c = ws.cell(row=avg_row, column=c_idx)
    c.value = f"=AVERAGE({c_let}{start_row}:{c_let}{avg_row-1})"
    c.number_format = '0.0%'
    c.font = Font(name=font_family, size=11, bold=True, color="1E3A8A")
    c.fill = overall_fill
    c.border = thin_border
    c.alignment = Alignment(horizontal='center')

    # Freeze Panes
    ws.freeze_panes = 'C6'

    wb.save(filename)
    print(f"Generated scheduled lesson gradebook: {filename}")

if __name__ == "__main__":
    create_gradebook_excel("Teacher_Production_Gradebook.xlsx")
